import os, re, io, json, hmac, base64, hashlib, secrets, asyncio
from datetime import datetime, date, timedelta
from pathlib import Path

from fastapi import FastAPI, Request, Response, Depends, HTTPException, UploadFile, File
from fastapi.responses import HTMLResponse, StreamingResponse
from sqlalchemy import create_engine, Column, Integer, String, Date, DateTime, Float, Text, ForeignKey, Boolean, func, inspect, text
from sqlalchemy.orm import declarative_base, sessionmaker, relationship, Session
from openpyxl import load_workbook, Workbook

BASE_DIR=Path(__file__).resolve().parent
APP_VERSION="2026.09.30.4"
APP_SECRET=os.getenv("APP_SECRET") or secrets.token_urlsafe(48)
DATABASE_URL=os.getenv("DATABASE_URL","sqlite:///./pedidoflow.db")
if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL="postgresql+psycopg://"+DATABASE_URL[len("postgres://"):]
elif DATABASE_URL.startswith("postgresql://"):
    DATABASE_URL="postgresql+psycopg://"+DATABASE_URL[len("postgresql://"):]

engine=create_engine(DATABASE_URL,connect_args={"check_same_thread":False} if DATABASE_URL.startswith("sqlite") else {},pool_pre_ping=True)
SessionLocal=sessionmaker(bind=engine,autoflush=False,autocommit=False)
Base=declarative_base()

class Organization(Base):
    __tablename__="organizations"
    id=Column(Integer,primary_key=True)
    name=Column(String(160),nullable=False)

class User(Base):
    __tablename__="users"
    id=Column(Integer,primary_key=True)
    organization_id=Column(Integer,ForeignKey("organizations.id"),nullable=False)
    name=Column(String(120),nullable=False)
    username=Column(String(80))
    email=Column(String(180),nullable=False)
    password_hash=Column(String(300),nullable=False)
    role=Column(String(20),default="buyer")
    active=Column(Boolean,default=True)
    must_change_password=Column(Boolean)

class Supplier(Base):
    __tablename__="suppliers"
    id=Column(Integer,primary_key=True)
    organization_id=Column(Integer,ForeignKey("organizations.id"),nullable=False)
    legal_name=Column(String(240),nullable=False)
    trade_name=Column(String(180))
    tax_id=Column(String(40))
    phone=Column(String(60))
    email=Column(String(180))

class PurchaseOrder(Base):
    __tablename__="purchase_orders"
    id=Column(Integer,primary_key=True)
    organization_id=Column(Integer,ForeignKey("organizations.id"),nullable=False)
    number=Column(String(30),nullable=False)
    supplier_id=Column(Integer,ForeignKey("suppliers.id"))
    buyer_id=Column(Integer,ForeignKey("users.id"))
    buyer_name=Column(String(120))
    issued_at=Column(Date)
    system_due_date=Column(Date)
    sent_at=Column(Date)
    priority=Column(String(20),default="Normal")
    entry_notes=Column(Text)
    delivery_type=Column(String(20))
    proposal_due_date=Column(Date)
    supplier_due_date=Column(Date)
    supplier_status=Column(String(100))
    next_follow_up=Column(Date)
    tracking_notes=Column(Text)
    source_file=Column(String(260))
    source_imported_at=Column(DateTime)
    created_at=Column(DateTime,default=datetime.utcnow)
    last_modified_at=Column(DateTime)
    last_modified_by=Column(String(120))
    updated_at=Column(DateTime,default=datetime.utcnow,onupdate=datetime.utcnow)
    supplier=relationship("Supplier")
    items=relationship("OrderItem",cascade="all, delete-orphan",back_populates="order")

class OrderItem(Base):
    __tablename__="order_items"
    id=Column(Integer,primary_key=True)
    order_id=Column(Integer,ForeignKey("purchase_orders.id"),nullable=False)
    erp_code=Column(String(60))
    description=Column(String(500),nullable=False)
    unit=Column(String(20))
    quantity=Column(Float,default=0)
    quantity_delivered=Column(Float,default=0)
    total_value=Column(Float,default=0)
    order=relationship("PurchaseOrder",back_populates="items")

class ImportBatch(Base):
    __tablename__="import_batches"
    id=Column(Integer,primary_key=True)
    organization_id=Column(Integer,ForeignKey("organizations.id"),nullable=False)
    filename=Column(String(260),nullable=False)
    imported_at=Column(DateTime,default=datetime.utcnow)
    imported_by=Column(String(120))
    orders_in_file=Column(Integer,default=0)
    new_orders=Column(Integer,default=0)
    skipped_orders=Column(Integer,default=0)
    item_rows=Column(Integer,default=0)

Base.metadata.create_all(engine)

def ensure_columns():
    additions={
        "delivery_type":"VARCHAR(20)",
        "proposal_due_date":"DATE",
        "source_file":"VARCHAR(260)",
        "source_imported_at":"TIMESTAMP",
        "created_at":"TIMESTAMP",
        "last_modified_at":"TIMESTAMP",
        "last_modified_by":"VARCHAR(120)",
    }
    existing={c["name"] for c in inspect(engine).get_columns("purchase_orders")}
    user_existing={c["name"] for c in inspect(engine).get_columns("users")}
    with engine.begin() as conn:
        for name,ddl in additions.items():
            if name not in existing:
                conn.execute(text(f"ALTER TABLE purchase_orders ADD COLUMN {name} {ddl}"))
        if "username" not in user_existing:
            conn.execute(text("ALTER TABLE users ADD COLUMN username VARCHAR(80)"))
        if "must_change_password" not in user_existing:
            conn.execute(text("ALTER TABLE users ADD COLUMN must_change_password BOOLEAN"))
        conn.execute(text("UPDATE purchase_orders SET created_at = COALESCE(created_at, updated_at, CURRENT_TIMESTAMP)"))
        if engine.dialect.name=="sqlite":
            conn.execute(text("UPDATE users SET username = LOWER(SUBSTR(email,1,INSTR(email,'@')-1)) WHERE (username IS NULL OR username='') AND email LIKE '%@%'"))
        else:
            conn.execute(text("UPDATE users SET username = LOWER(SPLIT_PART(email,'@',1)) WHERE (username IS NULL OR username='') AND email LIKE '%@%'"))
        admin_row=conn.execute(text("SELECT id, username FROM users WHERE LOWER(COALESCE(username,''))='weverson' LIMIT 1")).fetchone()
        if not admin_row:
            legacy=conn.execute(text("SELECT id FROM users WHERE role='admin' ORDER BY id LIMIT 1")).fetchone()
            if legacy:
                conn.execute(text("UPDATE users SET username='weverson', role='admin' WHERE id=:id"),{"id":legacy[0]})
        conn.execute(text("UPDATE users SET role='buyer' WHERE LOWER(COALESCE(username,''))<>'weverson'"))
        conn.execute(text("UPDATE users SET role='admin' WHERE LOWER(COALESCE(username,''))='weverson'"))
ensure_columns()

app=FastAPI(title="PedidoFlow Online")
subscribers={}

def dbdep():
    db=SessionLocal()
    try: yield db
    finally: db.close()

def hpw(p):
    salt=secrets.token_bytes(16)
    dk=hashlib.pbkdf2_hmac("sha256",p.encode(),salt,180000)
    return "pbkdf2$"+base64.b64encode(salt).decode()+"$"+base64.b64encode(dk).decode()

def vpw(p,s):
    try:
        _,a,b=s.split("$",2)
        salt=base64.b64decode(a); expected=base64.b64decode(b)
        got=hashlib.pbkdf2_hmac("sha256",p.encode(),salt,180000)
        return hmac.compare_digest(got,expected)
    except: return False

def ensure_default_admin():
    db=SessionLocal()
    try:
        admin_user=db.query(User).filter(func.lower(User.username)=="weverson").first()
        if admin_user:
            if admin_user.role!="admin":
                admin_user.role="admin"
                db.commit()
            return
        org=db.query(Organization).order_by(Organization.id).first()
        if not org:
            org=Organization(name="EMPAT")
            db.add(org);db.flush()
        admin_user=User(
            organization_id=org.id,
            name="WEVERSON",
            username="weverson",
            email="weverson@pedidoflow.local",
            password_hash=hpw("123456"),
            role="admin",
            active=True,
            must_change_password=False
        )
        db.add(admin_user);db.commit()
    finally:
        db.close()

ensure_default_admin()

def initialize_temp_passwords():
    db=SessionLocal()
    try:
        pending=db.query(User).filter(User.must_change_password==None).all()
        for u in pending:
            u.password_hash=hpw("123456")
            u.must_change_password=True
        if pending:
            db.commit()
    finally:
        db.close()

initialize_temp_passwords()

def token(uid,oid):
    body=base64.urlsafe_b64encode(json.dumps({"uid":uid,"oid":oid,"ver":APP_VERSION,"exp":int((datetime.utcnow()+timedelta(days=7)).timestamp())},separators=(",",":")).encode()).decode().rstrip("=")
    sig=hmac.new(APP_SECRET.encode(),body.encode(),hashlib.sha256).hexdigest()
    return body+"."+sig

def untoken(t):
    try:
        body,sig=t.split(".",1)
        if not hmac.compare_digest(sig,hmac.new(APP_SECRET.encode(),body.encode(),hashlib.sha256).hexdigest()): return None
        p=json.loads(base64.urlsafe_b64decode(body+"="*(-len(body)%4)))
        if p["exp"]<int(datetime.utcnow().timestamp()): return None
        if p.get("ver")!=APP_VERSION: return None
        return p
    except: return None

def current(req:Request,db:Session=Depends(dbdep)):
    p=untoken(req.cookies.get("session",""))
    if not p: raise HTTPException(401,"Faça login")
    u=db.query(User).filter(User.id==p["uid"],User.organization_id==p["oid"],User.active==True).first()
    if not u: raise HTTPException(401,"Sessão inválida")
    return u

def admin(u=Depends(current)):
    if u.role!="admin": raise HTTPException(403,"Acesso administrativo")
    return u

def dt(v):
    if not v:return None
    if isinstance(v,datetime):return v.date()
    if isinstance(v,date):return v
    s=str(v).strip()
    for f in ("%d/%m/%Y","%Y-%m-%d","%d/%m/%y"):
        try:return datetime.strptime(s[:10],f).date()
        except: pass
    return None

def num(v):
    if v in (None,""):return 0.0
    if isinstance(v,(int,float)):return float(v)
    s=str(v).strip().replace(".","").replace(",",".")
    try:return float(s)
    except:return 0.0

def norm(s):
    return re.sub(r"\s+"," ",str(s or "").strip())

def workflow(o):
    has_owner=bool(o.buyer_name)
    complete=bool(o.buyer_name and o.sent_at and o.delivery_type)
    if complete:return "Preenchido"
    if has_owner or o.sent_at or o.delivery_type:return "Pendente"
    return "Novo"

def situation(o):
    q=sum(i.quantity or 0 for i in o.items); d=sum(i.quantity_delivered or 0 for i in o.items)
    today=date.today(); due=o.supplier_due_date or o.system_due_date
    if q and d>=q:return "Concluído"
    if o.supplier_due_date and o.supplier_due_date<today:return "Atrasado"
    if d>0 and d<q:return "Entrega parcial"
    if o.next_follow_up and o.next_follow_up<=today:return "Cobrar hoje"
    if due and due<=today+timedelta(days=3):return "Prazo próximo"
    return "Em acompanhamento"

def iso_dt(v):
    return v.isoformat(timespec="seconds") if v else None

def order_json(o):
    q=sum(i.quantity or 0 for i in o.items); d=sum(i.quantity_delivered or 0 for i in o.items)
    pct=round((d/q*100),1) if q else 0
    return {
      "id":o.id,"number":o.number,"workflow_status":workflow(o),"situation":situation(o),
      "buyer_name":o.buyer_name,"buyer_id":o.buyer_id,
      "sent_at":o.sent_at.isoformat() if o.sent_at else None,"priority":o.priority or "Normal",
      "entry_notes":o.entry_notes or "","delivery_type":o.delivery_type or "",
      "system_due_date":o.system_due_date.isoformat() if o.system_due_date else None,
      "proposal_due_date":o.proposal_due_date.isoformat() if o.proposal_due_date else None,
      "supplier_due_date":o.supplier_due_date.isoformat() if o.supplier_due_date else None,
      "supplier_status":o.supplier_status or "","next_follow_up":o.next_follow_up.isoformat() if o.next_follow_up else None,
      "tracking_notes":o.tracking_notes or "","source_file":o.source_file or "","source_imported_at":iso_dt(o.source_imported_at),"created_at":iso_dt(o.created_at),
      "last_modified_at":iso_dt(o.last_modified_at or o.updated_at or o.created_at),
      "last_modified_by":o.last_modified_by or "",
      "quantity":q,"delivered":d,"percent":pct,
      "supplier":{"name":(o.supplier.trade_name or o.supplier.legal_name) if o.supplier else "",
                  "legal_name":o.supplier.legal_name if o.supplier else "",
                  "phone":o.supplier.phone if o.supplier else "","email":o.supplier.email if o.supplier else ""},
      "items":[{"code":i.erp_code,"description":i.description,"unit":i.unit,"quantity":i.quantity,
                "delivered":i.quantity_delivered,"total_value":i.total_value} for i in o.items]
    }

def publish(org_id,event_type,payload=None):
    msg=json.dumps({"type":event_type,"payload":payload or {},"at":datetime.utcnow().isoformat()},ensure_ascii=False)
    for q in list(subscribers.get(org_id,set())):
        try:q.put_nowait(msg)
        except:pass

async def event_stream(org_id):
    q=asyncio.Queue(maxsize=50)
    subscribers.setdefault(org_id,set()).add(q)
    try:
        yield "event: ready\ndata: {}\n\n"
        while True:
            try:
                msg=await asyncio.wait_for(q.get(),timeout=20)
                yield f"data: {msg}\n\n"
            except asyncio.TimeoutError:
                yield ": ping\n\n"
    finally:
        subscribers.get(org_id,set()).discard(q)

@app.get("/healthz")
def health(): return {"ok":True,"service":"PedidoFlow Online","database":"postgres" if DATABASE_URL.startswith("postgresql") else "sqlite"}

@app.get("/api/events")
async def events(u=Depends(current)):
    return StreamingResponse(event_stream(u.organization_id),media_type="text/event-stream",headers={"Cache-Control":"no-cache","X-Accel-Buffering":"no"})

@app.get("/api/bootstrap")
def bootstrap(req:Request,res:Response,db:Session=Depends(dbdep)):
    cookie=req.cookies.get("session","")
    initialized=db.query(User).count()>0
    p=untoken(cookie) if cookie else None
    user=None
    update_required=False
    if p:
        u=db.query(User).filter(User.id==p["uid"],User.active==True).first()
        if u:
            user={"id":u.id,"name":u.name,"username":u.username or u.name,"role":u.role,"must_change_password":bool(u.must_change_password)}
    elif cookie:
        update_required=True
        res.delete_cookie("session")
    return {"initialized":initialized,"user":user,"update_required":update_required,"version":APP_VERSION}

def clean_username(value):
    username=norm(value).lower()
    if len(username)<3 or len(username)>40:
        raise HTTPException(400,"Usuário deve ter entre 3 e 40 caracteres")
    if not re.fullmatch(r"[a-z0-9._-]+",username):
        raise HTTPException(400,"Use somente letras, números, ponto, hífen ou underline no usuário")
    return username

def default_display_name(username):
    return username.replace("."," ").replace("_"," ").replace("-"," ").strip().upper()

@app.post("/api/login")
async def login(data:dict,res:Response,db:Session=Depends(dbdep)):
    username=clean_username(data.get("username"))
    password=data.get("password") or ""
    if not password:
        raise HTTPException(400,"Informe usuário e senha")

    u=db.query(User).filter(func.lower(User.username)==username,User.active==True).first()

    if u:
        if not vpw(password,u.password_hash):
            raise HTTPException(401,"Usuário ou senha inválidos")
        expected_role="admin" if username=="weverson" else "buyer"
        if u.role!=expected_role:
            u.role=expected_role
            db.commit()
    else:
        if password!="123456":
            raise HTTPException(401,"Usuário ou senha inválidos")

        if username=="weverson":
            org=db.query(Organization).order_by(Organization.id).first()
            if not org:
                org=Organization(name="EMPAT")
                db.add(org);db.flush()
            u=User(
                organization_id=org.id,
                name="WEVERSON",
                username="weverson",
                email="weverson@pedidoflow.local",
                password_hash=hpw("123456"),
                role="admin",
                active=True,
                must_change_password=False
            )
        else:
            admin_user=db.query(User).filter(func.lower(User.username)=="weverson",User.active==True).first()
            if not admin_user:
                org=db.query(Organization).order_by(Organization.id).first()
                if not org:
                    org=Organization(name="EMPAT");db.add(org);db.flush()
                admin_user=User(organization_id=org.id,name="WEVERSON",username="weverson",email="weverson@pedidoflow.local",password_hash=hpw("123456"),role="admin",active=True,must_change_password=False)
                db.add(admin_user);db.flush()
            u=User(
                organization_id=admin_user.organization_id,
                name=default_display_name(username),
                username=username,
                email=f"{username}@pedidoflow.local",
                password_hash=hpw("123456"),
                role="buyer",
                active=True,
                must_change_password=False
            )
        db.add(u);db.commit();db.refresh(u)
        publish(u.organization_id,"user.created",{"name":u.name,"username":u.username,"role":u.role})

    res.set_cookie("session",token(u.id,u.organization_id),httponly=True,samesite="lax",secure=True,max_age=604800)
    return {"ok":True,"role":u.role,"name":u.name,"username":u.username,"created":bool(not db.query(User).filter(User.id==u.id).first() is None)}

@app.post("/api/change-password")
async def change_password(data:dict,u=Depends(current),db:Session=Depends(dbdep)):
    current_password=data.get("current_password") or ""
    new_password=data.get("new_password") or ""
    if not vpw(current_password,u.password_hash):
        raise HTTPException(400,"Senha atual incorreta")
    if len(new_password)<6:
        raise HTTPException(400,"A nova senha deve ter pelo menos 6 caracteres")
    if new_password=="123456":
        raise HTTPException(400,"Escolha uma senha diferente da senha inicial")
    if current_password==new_password:
        raise HTTPException(400,"A nova senha deve ser diferente da senha atual")
    u.password_hash=hpw(new_password)
    u.must_change_password=False
    db.commit()
    return {"ok":True}

@app.post("/api/logout")
def logout(res:Response):
    res.delete_cookie("session");return {"ok":True}

@app.get("/api/users")
def users(u=Depends(admin),db:Session=Depends(dbdep)):
    return [{"id":x.id,"name":x.name,"username":x.username,"role":x.role,"must_change_password":bool(x.must_change_password)} for x in db.query(User).filter(User.organization_id==u.organization_id).all()]

@app.post("/api/users")
async def create_user(data:dict,u=Depends(admin),db:Session=Depends(dbdep)):
    name=norm(data.get("name"))
    username=clean_username(data.get("username"))
    if username=="weverson":
        raise HTTPException(400,"Este usuário já está reservado")
    if db.query(User).filter(User.organization_id==u.organization_id,func.lower(User.username)==username).first():
        raise HTTPException(400,"Este usuário já existe")
    x=User(
        organization_id=u.organization_id,
        name=name or default_display_name(username),
        username=username,
        email=f"{username}@pedidoflow.local",
        password_hash=hpw("123456"),
        role="buyer",
        active=True,
        must_change_password=False
    )
    db.add(x);db.commit();db.refresh(x)
    publish(u.organization_id,"user.created",{"name":x.name,"username":x.username})
    return {"ok":True,"id":x.id,"name":x.name,"username":x.username}

@app.get("/api/orders")
def orders(search:str="",u=Depends(current),db:Session=Depends(dbdep)):
    q=db.query(PurchaseOrder).filter(PurchaseOrder.organization_id==u.organization_id)
    if u.role=="buyer":
        q=q.filter((PurchaseOrder.buyer_id==u.id)|(PurchaseOrder.buyer_id==None))
    data=q.order_by(PurchaseOrder.created_at.desc(),PurchaseOrder.number.desc()).all()
    s=search.lower().strip()
    if s:
        data=[o for o in data if s in o.number.lower()
              or (o.supplier and s in ((o.supplier.legal_name or "")+" "+(o.supplier.trade_name or "")).lower())
              or any(s in (i.description or "").lower() or s in (i.erp_code or "").lower() for i in o.items)]
    return [order_json(o) for o in data[:1000]]

@app.get("/api/orders/{number}")
def one(number:str,u=Depends(current),db:Session=Depends(dbdep)):
    o=db.query(PurchaseOrder).filter(PurchaseOrder.organization_id==u.organization_id,PurchaseOrder.number==number.zfill(6)).first()
    if not o:raise HTTPException(404,"Pedido não encontrado")
    if u.role=="buyer" and o.buyer_id not in (None,u.id):raise HTTPException(403,"Pedido atribuído a outro comprador")
    return order_json(o)

@app.post("/api/orders/{number}/buyer-submit")
async def buyer_submit(number:str,data:dict,u=Depends(current),db:Session=Depends(dbdep)):
    if u.role not in ("buyer","admin"):raise HTTPException(403)
    o=db.query(PurchaseOrder).filter(PurchaseOrder.organization_id==u.organization_id,PurchaseOrder.number==number.zfill(6)).first()
    if not o:raise HTTPException(404,"Pedido não encontrado")
    if o.buyer_id not in (None,u.id) and u.role!="admin":raise HTTPException(403,"Pedido atribuído a outro comprador")
    sent=dt(data.get("sent_at")); delivery=(data.get("delivery_type") or "").upper()
    if not sent:raise HTTPException(400,"Informe a data de envio ao fornecedor")
    if delivery not in ("COLETA","ENTREGA"):raise HTTPException(400,"Selecione Coleta ou Entrega")
    o.buyer_id=u.id;o.buyer_name=u.name;o.sent_at=sent;o.delivery_type=delivery
    o.priority=data.get("priority") if data.get("priority") in ("Baixa","Normal","Alta","Urgente") else "Normal"
    o.entry_notes=norm(data.get("entry_notes"))
    o.proposal_due_date=dt(data.get("proposal_due_date")) if data.get("proposal_due_date") else None
    o.last_modified_at=datetime.utcnow();o.last_modified_by=u.name
    db.commit();db.refresh(o)
    publish(u.organization_id,"buyer.submitted",{"number":o.number,"buyer":u.name,"workflow_status":workflow(o)})
    return order_json(o)

@app.patch("/api/orders/{number}")
async def admin_update(number:str,data:dict,u=Depends(admin),db:Session=Depends(dbdep)):
    o=db.query(PurchaseOrder).filter(PurchaseOrder.organization_id==u.organization_id,PurchaseOrder.number==number.zfill(6)).first()
    if not o:raise HTTPException(404)
    for k in ("supplier_status","tracking_notes"):
        if k in data:setattr(o,k,norm(data[k]))
    for k in ("proposal_due_date","supplier_due_date","next_follow_up"):
        if k in data:setattr(o,k,dt(data[k]))
    if "delivery_type" in data:
        dv=(data.get("delivery_type") or "").upper()
        if dv in ("","COLETA","ENTREGA"):o.delivery_type=dv or None
    o.last_modified_at=datetime.utcnow();o.last_modified_by=u.name
    db.commit();db.refresh(o)
    publish(u.organization_id,"order.updated",{"number":o.number,"by":u.name})
    return order_json(o)

@app.post("/api/import/orders")
async def import_orders(file:UploadFile=File(...),u=Depends(admin),db:Session=Depends(dbdep)):
    filename=norm(file.filename) or "relatorio_pedidos.xlsx"
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(400,"Envie um arquivo .xlsx")
    raw=await file.read()
    wb=load_workbook(io.BytesIO(raw),data_only=True)
    ws=wb.active
    grouped={}
    for r in ws.iter_rows(min_row=4,values_only=True):
        pc=norm(r[0])
        if not pc or not re.fullmatch(r"\d{1,6}",pc):
            continue
        pc=pc.zfill(6)
        g=grouped.setdefault(pc,{"issued_at":dt(r[1]),"supplier":norm(r[9]) or "Fornecedor não informado","system_due_date":dt(r[10]),"items":[]})
        if not g["issued_at"]:g["issued_at"]=dt(r[1])
        if not g["system_due_date"]:g["system_due_date"]=dt(r[10])
        g["items"].append({
            "code":norm(r[2]),
            "description":norm(r[3]) or "Item",
            "unit":norm(r[4]),
            "quantity":num(r[5]),
            "total_value":num(r[7]),
            "delivered":num(r[11])
        })

    created=skipped=items=0
    now=datetime.utcnow()
    for pc,g in grouped.items():
        existing=db.query(PurchaseOrder).filter(
            PurchaseOrder.organization_id==u.organization_id,
            PurchaseOrder.number==pc
        ).first()
        if existing:
            skipped+=1
            continue

        sup=db.query(Supplier).filter(
            Supplier.organization_id==u.organization_id,
            func.lower(Supplier.legal_name)==g["supplier"].lower()
        ).first()
        if not sup:
            sup=Supplier(organization_id=u.organization_id,legal_name=g["supplier"])
            db.add(sup);db.flush()

        o=PurchaseOrder(
            organization_id=u.organization_id,
            number=pc,
            supplier_id=sup.id,
            issued_at=g["issued_at"],
            system_due_date=g["system_due_date"],
            created_at=now,
            source_file=filename,
            source_imported_at=now,
            priority="Normal"
        )
        db.add(o);db.flush()
        created+=1
        for it in g["items"]:
            db.add(OrderItem(
                order_id=o.id,
                erp_code=it["code"],
                description=it["description"],
                unit=it["unit"],
                quantity=it["quantity"],
                total_value=it["total_value"],
                quantity_delivered=it["delivered"]
            ))
            items+=1

    batch=ImportBatch(
        organization_id=u.organization_id,
        filename=filename,
        imported_at=now,
        imported_by=u.name,
        orders_in_file=len(grouped),
        new_orders=created,
        skipped_orders=skipped,
        item_rows=items
    )
    db.add(batch)
    db.commit()
    publish(u.organization_id,"import.completed",{"filename":filename,"orders":len(grouped),"created":created,"skipped":skipped,"items":items})
    return {"ok":True,"filename":filename,"orders":len(grouped),"created":created,"skipped":skipped,"items":items}

@app.get("/api/import/history")
def import_history(u=Depends(admin),db:Session=Depends(dbdep)):
    rows=db.query(ImportBatch).filter(ImportBatch.organization_id==u.organization_id).order_by(ImportBatch.imported_at.desc()).limit(50).all()
    return [{
        "id":x.id,
        "filename":x.filename,
        "imported_at":iso_dt(x.imported_at),
        "imported_by":x.imported_by or "",
        "orders_in_file":x.orders_in_file or 0,
        "new_orders":x.new_orders or 0,
        "skipped_orders":x.skipped_orders or 0,
        "item_rows":x.item_rows or 0
    } for x in rows]

@app.post("/api/import/suppliers")
async def import_suppliers(file:UploadFile=File(...),u=Depends(admin),db:Session=Depends(dbdep)):
    if not (file.filename or "").lower().endswith(".xlsx"):raise HTTPException(400,"Envie um arquivo .xlsx")
    raw=await file.read();wb=load_workbook(io.BytesIO(raw),data_only=True);ws=wb.active
    headers=[norm(c.value).lower() for c in ws[1]]
    def ix(*names):
        for n in names:
            for j,h in enumerate(headers):
                if n in h:return j
        return None
    ni=ix("razão social","razao social","fornecedor");ti=ix("fantasia");ci=ix("cnpj");ddi=ix("ddd");phi=ix("telefone")
    n=0
    if ni is None:raise HTTPException(400,"Não encontrei a coluna Razão Social/Fornecedor")
    for row in ws.iter_rows(min_row=2,values_only=True):
        name=norm(row[ni])
        if not name:continue
        s=db.query(Supplier).filter(Supplier.organization_id==u.organization_id,func.lower(Supplier.legal_name)==name.lower()).first()
        if not s:s=Supplier(organization_id=u.organization_id,legal_name=name);db.add(s)
        if ti is not None:s.trade_name=norm(row[ti])
        if ci is not None:s.tax_id=norm(row[ci])
        phone=(norm(row[ddi]) if ddi is not None else "")+(norm(row[phi]) if phi is not None else "")
        if phone:s.phone=phone
        n+=1
    db.commit();publish(u.organization_id,"suppliers.imported",{"rows":n})
    return {"ok":True,"rows":n}

@app.get("/api/export.xlsx")
def export_xlsx(u=Depends(admin),db:Session=Depends(dbdep)):
    wb=Workbook();ws=wb.active;ws.title="ENTRADA_COMPRADORES"
    ws.append(["Nº Pedido","Comprador responsável","Data envio","Prioridade","Tipo entrega","Previsão entrega proposta","Observação","Fornecedor","Descrição dos itens","Prazo confirmado","Status fornecedor","Próxima cobrança","Data alteração","Alterado por","Situação preenchimento"])
    for o in db.query(PurchaseOrder).filter(PurchaseOrder.organization_id==u.organization_id).order_by(PurchaseOrder.number).all():
        desc=" | ".join(i.description for i in o.items)
        ws.append([o.number,o.buyer_name,o.sent_at,o.priority,o.delivery_type,o.proposal_due_date,o.entry_notes,(o.supplier.trade_name or o.supplier.legal_name) if o.supplier else "",desc,o.supplier_due_date,o.supplier_status,o.next_follow_up,o.last_modified_at or o.updated_at,o.last_modified_by,workflow(o)])
    out=io.BytesIO();wb.save(out);out.seek(0)
    return StreamingResponse(out,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=PedidoFlow_Export.xlsx"})

@app.get("/",response_class=HTMLResponse)
def root(): return (BASE_DIR/"login.html").read_text(encoding="utf-8")

@app.get("/admin",response_class=HTMLResponse)
def admin_page(): return (BASE_DIR/"admin.html").read_text(encoding="utf-8")

@app.get("/painel",response_class=HTMLResponse)
def buyer_page(): return (BASE_DIR/"buyer.html").read_text(encoding="utf-8")

@app.get("/comprador",response_class=HTMLResponse)
def buyer_legacy(): return (BASE_DIR/"login.html").read_text(encoding="utf-8")
